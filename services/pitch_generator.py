"""
Tailored Medical Pitch Generator and PowerPoint (.pptx) Compiler.
Synthesizes company workforce research and candidate policy clauses from Qdrant
into a tailored 3-5 slide commercial insurance presentation deck.
"""
import os
import re
import json
import logging
from pathlib import Path
from typing import Dict, Any, List, Optional
from django.conf import settings

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

from services.llm_client import generate_structured
from services.policy_retrieval import search_policy

logger = logging.getLogger(__name__)


# ==================================================
# PITCH GENERATOR LOGIC
# ==================================================

def generate_medical_pitch(
    company_profile: Dict[str, Any],
    workforce_health: Dict[str, Any],
    document_ids: Optional[List[str]] = None
) -> Dict[str, Any]:
    """
    Generates a tailored 3-5 slide marketing pitch linking company workforce health exposures
    to candidate medical policy clauses retrieved from Qdrant, and compiles a .pptx presentation.
    
    Args:
        company_profile: Dict from research_company()
        workforce_health: Dict from formulate_workforce_health_exposures()
        document_ids: Optional list of policy document IDs to filter by in Qdrant
        
    Returns:
        Dict containing:
        - company_name: str
        - presentation_title: str
        - slides: List[Dict[str, Any]]
        - pptx_url: str
        - pptx_path: str
        - candidate_evidence_chunks: List[Dict[str, Any]]
    """
    company_name = company_profile.get("company_name", "Corporate Client")
    industry = company_profile.get("industry_sector") or company_profile.get("industry", "Corporate Commercial")
    size = company_profile.get("company_size", "Enterprise scale")
    footprint = company_profile.get("operational_footprint", [])
    queries = workforce_health.get("medical_search_queries", [])

    # 1. Candidate Retrieval from Qdrant
    candidate_chunks: List[Dict[str, Any]] = []
    seen_chunk_signatures = set()

    for q in queries:
        try:
            hits = search_policy(query=q, document_ids=document_ids, limit=4)
            for hit in hits:
                sig = (hit.get("document_name"), hit.get("page_number"), hit.get("text", "")[:80])
                if sig not in seen_chunk_signatures:
                    seen_chunk_signatures.add(sig)
                    candidate_chunks.append(hit)
        except Exception as e:
            logger.warning(f"Error querying Qdrant for query '{q}': {e}")

    # Fallback if no specific policy chunks were found
    if not candidate_chunks:
        try:
            candidate_chunks = search_policy(query="inpatient hospitalization ICU room rent air ambulance pre post medical expenses", limit=8)
        except Exception as e:
            logger.error(f"Fallback policy retrieval failed: {e}")

    # Validate that policy chunks are available
    if not candidate_chunks:
        raise ValueError(
            "No indexed policy documents found in the database. "
            "Please upload at least one insurance policy PDF brochure before generating the pitch deck."
        )

    # Globally sort candidate chunks by vector relevance score descending
    candidate_chunks.sort(key=lambda c: c.get("score", 0.0), reverse=True)

    # Format top candidate evidence for prompt (top 10 chunks)
    evidence_text_blocks = []
    for i, c in enumerate(candidate_chunks[:10], 1):
        doc_name = c.get("document_name", "Policy Document")
        page_num = c.get("page_number", 1)
        snippet = c.get("text", "").strip().replace("\n", " ")
        evidence_text_blocks.append(f"[{i}] Document: '{doc_name}' | Page {page_num}\nExcerpt: \"{snippet[:350]}\"")

    evidence_context = "\n\n".join(evidence_text_blocks) or "No policy evidence found in database."

    system_prompt = (
        "You are an elite Marsh corporate insurance advisory specialist.\n"
        "Your task is to synthesize a tailored 5-slide executive marketing presentation "
        "recommending comprehensive employee healthcare and medical insurance coverage for a corporate client.\n\n"
        "STRICT PITCH RULES:\n"
        "1. GROUNDING MANDATE: Every policy claim or coverage benefit mentioned MUST cite the exact "
        "policy document name and page number from the provided candidate policy excerpts.\n"
        "2. NO FABRICATION: Do not promise benefits (like unlimited overseas dental or cosmetic surgery) "
        "that are not mentioned in the policy excerpts.\n"
        "3. WORKFORCE TAILORING: Connect benefits directly to the client's actual workforce profile, "
        "operational plants, offices, and employee health exposures.\n"
        "4. MULTIPLE VERIFIED CLAIMS MANDATE: For underwriting claim audit thoroughness, you MUST generate "
        "at least 2 to 3 distinct verifiable policy claims on Slide 3 (Tailored Coverage Mapping), "
        "at least 2 to 3 distinct claims on Slide 4 (Policy Safeguards & Enhanced Terms), "
        "and 1 to 2 claims on Slide 5 (Final Recommended Policy). Total policy claims across the presentation MUST be between 5 and 8. "
        "Do not leave policy_claims empty or with only 1 claim on coverage slides.\n"
        "5. FULL DECK REQUIREMENT: You MUST generate all 5 slides (Slide 1, Slide 2, Slide 3, Slide 4, Slide 5). Do not omit or truncate any slides.\n"
        "6. MULTI-CARRIER COMPARATIVE MANDATE: When multiple policy documents are provided in the candidate excerpts (e.g. Care Supreme.pdf and HDFC Product Brochure.pdf), you MUST cite and compare terms from multiple policies across the deck rather than attributing all claims to a single carrier.\n"
        "7. CLIENT-TAILORED ADVISORY MANDATE: On Slide 1 and Slide 2 ('Why Choose Marsh'), do NOT output generic static boilerplates. You MUST tailor Marsh's value propositions and bullets specifically to the client's industry sector, operational footprint, and workforce health risks.\n"
        "8. Output format MUST strictly be valid JSON."
    )

    prompt = f"""Target Client: {company_name}
Industry Sector: {industry}
Workforce Profile: {size}
Operational Footprint: {footprint}
Workforce Healthcare Priorities: {json.dumps(workforce_health.get('priority_healthcare_exposures', []), indent=2)}

CANDIDATE POLICY EXCERPTS FROM UPLOADED DOCUMENTS:
{evidence_context}

Generate an executive 3 to 5 slide presentation pitch JSON conforming strictly to this format:
{{
    "company_name": "{company_name}",
    "presentation_title": "Comprehensive Workforce Healthcare & Risk Protection: Strategic Advisory for {company_name}",
    "presentation_subtitle": "Tailored Employee Medical Coverage, Critical Care Safeguards & Policy Terms",
    "slides": [
        {{
            "slide_number": 1,
            "slide_type": "company_overview",
            "title": "Executive Briefing: {company_name} Healthcare Strategy",
            "subtitle": "Aligning health protection to operational scale across facilities",
            "bullets": [
                "<Operational footprint across {company_name}'s specific facilities, plants, and office locations>",
                "<Strategic corporate mandate to shield {company_name}'s workforce against industry-specific health and hospitalization risks>",
                "<Implementation of tailored corporate medical coverage designed to safeguard productivity and talent retention>"
            ],
            "policy_claims": []
        }},
        {{
            "slide_number": 2,
            "slide_type": "why_choose_marsh",
            "title": "Why Choose Marsh: Enterprise Advisory Advantage",
            "subtitle": "Global brokerage scale, bespoke claims advocacy & proprietary risk analytics",
            "bullets": [
                "<Marsh's specialized carrier bargaining power and group premium leverage tailored specifically to {company_name}'s {industry} sector risk profile>",
                "<Dedicated claims advocacy and deep cashless hospital network access where {company_name}'s key facilities and offices operate>",
                "<Customized workforce health analytics and proactive wellness risk engineering targeting {company_name}'s specific workforce exposures>"
            ],
            "policy_claims": []
        }},
        {{
            "slide_number": 3,
            "slide_type": "tailored_coverage_mapping",
            "title": "Tailored Policy Coverage & Exposure Mapping",
            "subtitle": "Directly mapping workforce operational risks to verified policy clauses",
            "bullets": [
                "Comprehensive inpatient hospitalization and ICU room rent protection for acute events",
                "Emergency transit coverage for rapid medical escalation across distributed sites",
                "Advanced day-care surgical treatments without mandatory 24-hour hospitalization"
            ],
            "policy_claims": [
                {{
                    "claim_text": "<Specific policy coverage statement e.g. Inpatient care hospitalization covered for medically necessary treatment>",
                    "cited_policy": "<Exact document name from excerpts e.g. Care Supreme.pdf>",
                    "cited_page": <Page integer from excerpt e.g. 1>,
                    "relevance_to_workforce": "<Why this clause protects this company's workforce>"
                }},
                {{
                    "claim_text": "<Second specific coverage statement e.g. Emergency air ambulance coverage up to Sum Insured>",
                    "cited_policy": "<Exact document name from excerpts e.g. Care Supreme.pdf>",
                    "cited_page": <Page integer from excerpt e.g. 3>,
                    "relevance_to_workforce": "<Why this protects distributed or plant workforce>"
                }},
                {{
                    "claim_text": "<Third specific coverage statement e.g. Day care treatments requiring less than 24 hours hospitalization>",
                    "cited_policy": "<Exact document name from excerpts e.g. Care Supreme.pdf>",
                    "cited_page": <Page integer from excerpt e.g. 2>,
                    "relevance_to_workforce": "<Relevance to quick surgical turnaround and employee recovery>"
                }}
            ]
        }},
        {{
            "slide_number": 4,
            "slide_type": "policy_safeguards",
            "title": "Key Policy Safeguards & Comprehensive Terms",
            "subtitle": "Critical value enhancements under the evaluated medical policies",
            "bullets": [
                "Pre and post hospitalization medical expense coverage window for continuous recovery",
                "Cumulative bonus enhancements and restore benefits for dependent family members",
                "Modern robotic treatments, organ donor cover, and ICU room rent capping safeguards"
            ],
            "policy_claims": [
                {{
                    "claim_text": "<First safeguard benefit e.g. Pre-hospitalization medical expenses covered for 60 days directly related to care>",
                    "cited_policy": "<Exact document name from excerpts>",
                    "cited_page": <Page integer from excerpt>,
                    "relevance_to_workforce": "<Protection against preliminary diagnostic costs>"
                }},
                {{
                    "claim_text": "<Second safeguard benefit e.g. Post-hospitalization medical expenses covered for 180 days following discharge>",
                    "cited_policy": "<Exact document name from excerpts>",
                    "cited_page": <Page integer from excerpt>,
                    "relevance_to_workforce": "<Coverage for outpatient medication and recovery>"
                }},
                {{
                    "claim_text": "<Third safeguard benefit e.g. Cumulative bonus enhancement on claim-free policy renewals>",
                    "cited_policy": "<Exact document name from excerpts>",
                    "cited_page": <Page integer from excerpt>,
                    "relevance_to_workforce": "<Value compounding for employee family healthcare pool>"
                }}
            ]
        }},
        {{
            "slide_number": 5,
            "slide_type": "final_recommended_policy",
            "title": "Final Recommended Policy & Advisory Roadmap",
            "subtitle": "Selection of the primary recommended policy based on workforce exposure profile",
            "bullets": [
                "Primary Recommendation: Care Supreme (Care Health Insurance) selected for superior sub-limit flexibility and uncapped room-rent terms",
                "Structuring corporate sum insured tiers across executive, management, and operational plant workers",
                "Direct cashless hospital network integration and seamless digital claims onboarding"
            ],
            "policy_claims": [
                {{
                    "claim_text": "<Key recommended policy term e.g. Cashless hospitalization network access with comprehensive critical care eligibility>",
                    "cited_policy": "<Exact document name from excerpts>",
                    "cited_page": <Page integer from excerpt>,
                    "relevance_to_workforce": "<Primary decision criteria justifying policy selection>"
                }}
            ]
        }}
    ]
}}
"""

    try:
        pitch_json = generate_structured(prompt=prompt, system_prompt=system_prompt, max_tokens=3500)
        slides = pitch_json.get("slides", [])
        if not slides or len(slides) < 3:
            raise ValueError(f"LLM generated only {len(slides)} slides; expected at least 3 slides.")
    except Exception as e:
        logger.exception(f"Failed to generate structured pitch JSON for {company_name}: {e}")
        raise


    # 2. Build PowerPoint Presentation (.pptx)
    pptx_path, pptx_url = compile_powerpoint(pitch_json)
    pitch_json["pptx_path"] = str(pptx_path)
    pitch_json["pptx_url"] = pptx_url
    pitch_json["candidate_evidence_chunks"] = candidate_chunks

    return pitch_json


def generateMarketingPitch(
    company_profile: Dict[str, Any],
    workforce_health: Optional[Dict[str, Any]] = None,
    document_ids: Optional[List[str]] = None
) -> Dict[str, Any]:
    """
    Marsh Case Study Objective 1.3:
    Produces a 3–5 slide deck covering:
    - company overview
    - why choose Marsh
    - policy benefits mapped to exposures
    - one final recommended policy
    """
    if workforce_health is None:
        from services.workforce_health import formulate_workforce_health_exposures
        workforce_health = formulate_workforce_health_exposures(company_profile)
    return generate_medical_pitch(company_profile, workforce_health, document_ids)



# ==================================================
# POWERPOINT COMPILER (python-pptx)
# ==================================================

def compile_powerpoint(pitch_data: Dict[str, Any]) -> tuple[Path, str]:
    """
    Compiles structured pitch JSON into a branded, widescreen 16:9 PowerPoint (.pptx) file.
    """
    prs = Presentation()
    # 16:9 Widescreen (13.333 x 7.5 inches)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    company_name = pitch_data.get("company_name", "Corporate Client")
    company_slug = re.sub(r'[^a-zA-Z0-9_-]', '_', company_name.lower())
    presentations_dir = Path(settings.MEDIA_ROOT) / "presentations"
    presentations_dir.mkdir(parents=True, exist_ok=True)
    pptx_filename = f"{company_slug}_medical_pitch.pptx"
    output_path = presentations_dir / pptx_filename

    # Corporate Typography System
    FONT_HEADING = "Segoe UI"
    FONT_BODY = "Segoe UI"

    # Corporate Executive Color Palette
    COLOR_NAVY = RGBColor(10, 37, 64)        # Deep Navy #0A2540
    COLOR_BLUE = RGBColor(0, 102, 224)       # Executive Accent Blue #0066E0
    COLOR_SLATE = RGBColor(71, 85, 105)      # Slate Gray #475569
    COLOR_LIGHT_BG = RGBColor(248, 250, 252) # Off-white #F8FAFC
    COLOR_WHITE = RGBColor(255, 255, 255)
    COLOR_BORDER = RGBColor(226, 232, 240)   # Border #E2E8F0
    COLOR_ACCENT_BG = RGBColor(241, 245, 249)# Soft Slate Tint #F1F5F9
    COLOR_MUTED = RGBColor(148, 163, 184)    # Slate 400 #94A3B8
    COLOR_CIT_BLUE = RGBColor(37, 99, 235)   # Verification Citation #2563EB

    slides_data = pitch_data.get("slides", [])

    for s_idx, slide_info in enumerate(slides_data):
        slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank slide layout

        # Top Branded Header Bar
        header_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.3))
        header_bar.fill.solid()
        header_bar.fill.fore_color.rgb = COLOR_NAVY
        header_bar.line.fill.background()

        # Header Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.18), Inches(10.5), Inches(0.8))
        tf = title_box.text_frame
        tf.word_wrap = True
        p_title = tf.paragraphs[0]
        p_title.text = slide_info.get("title", f"Slide {s_idx + 1}")
        p_title.font.name = FONT_HEADING
        p_title.font.size = Pt(23)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_WHITE

        # Subtitle
        if slide_info.get("subtitle"):
            p_sub = tf.add_paragraph()
            p_sub.text = slide_info.get("subtitle")
            p_sub.font.name = FONT_BODY
            p_sub.font.size = Pt(13)
            p_sub.font.color.rgb = RGBColor(203, 213, 225)
            p_sub.space_before = Pt(4)

        # Brand Tag in Header
        brand_box = slide.shapes.add_textbox(Inches(10.8), Inches(0.25), Inches(2.0), Inches(0.7))
        btf = brand_box.text_frame
        bp = btf.paragraphs[0]
        bp.alignment = PP_ALIGN.RIGHT
        bp.text = "MARSH ADVISORY"
        bp.font.name = FONT_HEADING
        bp.font.size = Pt(11)
        bp.font.bold = True
        bp.font.color.rgb = RGBColor(56, 189, 248)

        bp2 = btf.add_paragraph()
        bp2.alignment = PP_ALIGN.RIGHT
        bp2.text = "Client Strategic Brief"
        bp2.font.name = FONT_BODY
        bp2.font.size = Pt(9.5)
        bp2.font.color.rgb = COLOR_MUTED

        # Slide Body Layout
        claims = slide_info.get("policy_claims", [])

        if claims:
            # 2-Column Split: Left = Strategy Bullets, Right = Grounded Policy Evidence Cards
            # Left Card (Bullets)
            left_card = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
            left_card.fill.solid()
            left_card.fill.fore_color.rgb = COLOR_LIGHT_BG
            left_card.line.color.rgb = COLOR_BORDER
            left_card.line.width = Pt(1)

            # Left Text
            tb_left = slide.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.2), Inches(4.8))
            tf_l = tb_left.text_frame
            tf_l.word_wrap = True
            tf_l.margin_left = Inches(0.1)
            tf_l.margin_right = Inches(0.1)
            tf_l.margin_top = Inches(0.1)

            p_l_head = tf_l.paragraphs[0]
            p_l_head.text = "Strategic Workforce Protections"
            p_l_head.font.name = FONT_HEADING
            p_l_head.font.bold = True
            p_l_head.font.size = Pt(15)
            p_l_head.font.color.rgb = COLOR_NAVY

            for bullet in slide_info.get("bullets", []):
                p_b = tf_l.add_paragraph()
                p_b.text = f"•  {bullet}"
                p_b.font.name = FONT_BODY
                p_b.font.size = Pt(13.5)
                p_b.font.color.rgb = COLOR_SLATE
                p_b.space_before = Pt(14)

            # Right Card (Grounded Policy Claims)
            right_card = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2))
            right_card.fill.solid()
            right_card.fill.fore_color.rgb = COLOR_ACCENT_BG
            right_card.line.color.rgb = COLOR_BORDER
            right_card.line.width = Pt(1)

            tb_right = slide.shapes.add_textbox(Inches(7.0), Inches(1.8), Inches(5.3), Inches(4.8))
            tf_r = tb_right.text_frame
            tf_r.word_wrap = True
            tf_r.margin_left = Inches(0.1)
            tf_r.margin_right = Inches(0.1)
            tf_r.margin_top = Inches(0.1)

            p_r_head = tf_r.paragraphs[0]
            p_r_head.text = "Grounded Policy Terms & Page Citations"
            p_r_head.font.name = FONT_HEADING
            p_r_head.font.bold = True
            p_r_head.font.size = Pt(15)
            p_r_head.font.color.rgb = COLOR_NAVY

            for cl in claims[:3]:
                p_c = tf_r.add_paragraph()
                p_c.text = f"✓  {cl.get('claim_text')}"
                p_c.font.name = FONT_BODY
                p_c.font.bold = True
                p_c.font.size = Pt(13)
                p_c.font.color.rgb = COLOR_NAVY
                p_c.space_before = Pt(12)

                p_src = tf_r.add_paragraph()
                p_src.text = f"   [Source: {cl.get('cited_policy')}, Page {cl.get('cited_page', 1)}]"
                p_src.font.name = FONT_BODY
                p_src.font.size = Pt(10.5)
                p_src.font.bold = True
                p_src.font.color.rgb = COLOR_CIT_BLUE
                p_src.space_before = Pt(3)

                if cl.get("relevance_to_workforce"):
                    p_rel = tf_r.add_paragraph()
                    p_rel.text = f"   Benefit: {cl.get('relevance_to_workforce')}"
                    p_rel.font.name = FONT_BODY
                    p_rel.font.size = Pt(11)
                    p_rel.font.color.rgb = COLOR_SLATE
                    p_rel.space_before = Pt(3)

        else:
            # Full-Width Strategic Layout
            body_card = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.6), Inches(11.733), Inches(5.2))
            body_card.fill.solid()
            body_card.fill.fore_color.rgb = COLOR_LIGHT_BG
            body_card.line.color.rgb = COLOR_BORDER
            body_card.line.width = Pt(1)

            tb = slide.shapes.add_textbox(Inches(1.2), Inches(1.85), Inches(10.9), Inches(4.7))
            tf = tb.text_frame
            tf.word_wrap = True
            tf.margin_left = Inches(0.15)
            tf.margin_right = Inches(0.15)
            tf.margin_top = Inches(0.15)

            bullets = slide_info.get("bullets", [])
            for i, b in enumerate(bullets):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.text = f"•  {b}"
                p.font.name = FONT_BODY
                p.font.size = Pt(17.5)
                p.font.color.rgb = COLOR_NAVY
                p.space_before = Pt(22)

        # Footer Bar
        footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.9), Inches(11.733), Inches(0.4))
        ftf = footer_box.text_frame
        fp = ftf.paragraphs[0]
        fp.text = f"Marsh Commercial Advisory | Client Confidential: {company_name} | Slide {s_idx + 1} of {len(slides_data)}"
        fp.font.name = FONT_BODY
        fp.font.size = Pt(9.5)
        fp.font.color.rgb = COLOR_MUTED

    prs.save(str(output_path))
    pptx_url = f"{settings.MEDIA_URL}presentations/{pptx_filename}"
    return output_path, pptx_url





